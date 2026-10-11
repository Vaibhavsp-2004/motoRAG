"""
chunk_transcripts.py

Reads motorcycle_review_video_tracker.xlsx, finds every transcript marked
'done', splits each transcript into overlapping word-based chunks, and
writes them all out to a single chunks.jsonl file — one JSON object per
chunk, ready to be fed into an embedding model in the next pipeline step.

Run this on your own machine, in the same folder as the transcripts/
directory created by collect_transcripts.py.

Setup (one time):
    pip install openpyxl

Usage:
    python chunk_transcripts.py motorcycle_review_video_tracker.xlsx

Output:
    - chunks.jsonl   (one JSON object per line, one line per chunk)
"""

import sys
import json
from pathlib import Path

import openpyxl

# --- Chunking parameters ---
# 130 words keeps each chunk focused on roughly one topic (e.g. just seat
# height/ergonomics, or just the engine), so a specific fact doesn't get
# diluted by averaging its vector together with five unrelated topics.
CHUNK_SIZE_WORDS = 130

# Overlap re-includes the tail of one chunk at the start of the next, so a
# sentence or idea that happens to fall right on a chunk boundary doesn't
# get orphaned/split in a way that hurts retrieval.
OVERLAP_WORDS = 30

OUTPUT_PATH = Path("chunks.jsonl")


def load_motorcycle_metadata(ws) -> dict:
    """Build motorcycle_id -> {brand, model, variant, body_type, segment_group}."""
    headers = [cell.value for cell in ws[1]]
    col = {name: idx for idx, name in enumerate(headers)}
    meta = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        mid = row[col["motorcycle_id"]]
        if not mid:
            continue
        meta[mid] = {
            "brand": row[col["brand"]],
            "model": row[col["model"]],
            "variant": row[col["variant"]],
            "body_type": row[col["body_type"]],
            "segment_group": row[col["segment_group"]],
        }
    return meta


def chunk_text(text: str, chunk_size: int, overlap: int) -> list[str]:
    """Split text into overlapping chunks of `chunk_size` words."""
    words = text.split()
    if not words:
        return []

    chunks = []
    step = chunk_size - overlap
    if step <= 0:
        raise ValueError("overlap must be smaller than chunk_size")

    start = 0
    while start < len(words):
        chunk_words = words[start:start + chunk_size]
        if not chunk_words:
            break
        chunks.append(" ".join(chunk_words))
        if start + chunk_size >= len(words):
            break
        start += step

    return chunks


def main():
    if len(sys.argv) != 2:
        print("Usage: python chunk_transcripts.py <tracker.xlsx>")
        sys.exit(1)

    xlsx_path = Path(sys.argv[1])
    if not xlsx_path.exists():
        print(f"File not found: {xlsx_path}")
        sys.exit(1)

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws_videos = wb["videos"]
    ws_motos = wb["motorcycles"]

    moto_meta = load_motorcycle_metadata(ws_motos)

    headers = [cell.value for cell in ws_videos[1]]
    col = {name: idx for idx, name in enumerate(headers)}

    required = ["motorcycle_id", "youtube_url", "transcript_status", "transcript_path"]
    missing = [c for c in required if c not in col]
    if missing:
        print(f"Missing expected columns in 'videos' sheet: {missing}")
        sys.exit(1)

    total_transcripts = 0
    total_chunks = 0
    skipped_not_done = 0
    skipped_missing_file = 0

    with OUTPUT_PATH.open("w", encoding="utf-8") as out:
        for row in ws_videos.iter_rows(min_row=2, values_only=True):
            mid = row[col["motorcycle_id"]]
            url = row[col["youtube_url"]]
            status = row[col["transcript_status"]]
            transcript_path = row[col["transcript_path"]]

            if not url:
                continue  # no video collected for this motorcycle

            if status != "done" or not transcript_path:
                skipped_not_done += 1
                continue

            path = Path(transcript_path)
            if not path.exists():
                print(f"[{mid}] transcript file missing on disk: {transcript_path}")
                skipped_missing_file += 1
                continue

            text = path.read_text(encoding="utf-8", errors="ignore")
            chunks = chunk_text(text, CHUNK_SIZE_WORDS, OVERLAP_WORDS)

            meta = moto_meta.get(mid, {})
            bike_label = " ".join(str(x) for x in [
                meta.get("brand"), meta.get("model"), meta.get("variant")
            ] if x).strip()

            for i, chunk in enumerate(chunks):
                # embedding_text carries extra bike-identity context so the
                # vector for this chunk isn't just "generic seat height talk" -
                # it's clearly anchored to this specific motorcycle. The
                # original 'text' field stays clean for display/citation.
                embedding_text = f"{bike_label} review: {chunk}" if bike_label else chunk
                record = {
                    "chunk_id": f"{mid}_{i:03d}",
                    "motorcycle_id": mid,
                    "brand": meta.get("brand"),
                    "model": meta.get("model"),
                    "variant": meta.get("variant"),
                    "body_type": meta.get("body_type"),
                    "segment_group": meta.get("segment_group"),
                    "youtube_url": url,
                    "chunk_index": i,
                    "word_count": len(chunk.split()),
                    "text": chunk,
                    "embedding_text": embedding_text,
                }
                out.write(json.dumps(record, ensure_ascii=False) + "\n")
                total_chunks += 1

            total_transcripts += 1

    print("\n--- Summary ---")
    print(f"Transcripts chunked:        {total_transcripts}")
    print(f"Total chunks written:       {total_chunks}")
    if total_transcripts:
        print(f"Avg chunks per transcript:  {total_chunks / total_transcripts:.1f}")
    print(f"Skipped (not done yet):     {skipped_not_done}")
    print(f"Skipped (file missing):     {skipped_missing_file}")
    print(f"Output written to:          {OUTPUT_PATH.resolve()}")


if __name__ == "__main__":
    main()