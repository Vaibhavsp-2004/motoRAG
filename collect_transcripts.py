"""
collect_transcripts.py

Reads motorcycle_review_video_tracker.xlsx, pulls every youtube_url from the
'videos' sheet, downloads the captions/subtitles for each video via yt-dlp,
cleans them into plain-text transcripts, and writes the results back into
the workbook (transcript_status, transcript_path).

Run this on your own machine (not in a sandbox) since it needs real internet
access to YouTube.

Setup (one time):
    pip install yt-dlp openpyxl

Usage:
    python collect_transcripts.py motorcycle_review_video_tracker.xlsx

Output:
    - transcripts/<motorcycle_id>.txt   (one plain-text transcript per video)
    - the xlsx is updated in place with transcript_status + transcript_path
"""

import sys
import re
import subprocess
import time
from pathlib import Path

import openpyxl

TRANSCRIPTS_DIR = Path("transcripts")
LANG = "en"                 # caption language to try first
SLEEP_BETWEEN_VIDEOS = 6    # seconds, be polite to YouTube
YTDLP_TIMEOUT = 120         # seconds per video
MAX_RETRIES = 3
RATE_LIMIT_BACKOFF = 45     # seconds to wait after a 429 before retrying

# Set to "chrome", "edge", or "firefox" to reuse that browser's YouTube login
# cookies. Authenticated requests get a much higher rate limit than anonymous
# ones. Close the browser fully before running if you get a "database is
# locked" error. Leave as None to run anonymously.
COOKIES_FROM_BROWSER = None

BATCH_SIZE = 15             # proactively pause after this many videos
BATCH_PAUSE = 60            # seconds to pause between batches
CONSECUTIVE_FAIL_LIMIT = 3  # rate-limit failures in a row before a long cooldown
COOLDOWN_SECONDS = 300      # 5 minutes


def extract_video_id(url: str) -> str | None:
    """Pull the 11-char YouTube video ID out of a watch URL."""
    if not url:
        return None
    match = re.search(r"(?:v=|youtu\.be/)([A-Za-z0-9_-]{11})", url)
    return match.group(1) if match else None


def vtt_to_plain_text(vtt_path: Path) -> str:
    """Strip WEBVTT timing/cue markup down to plain, deduplicated text."""
    lines = vtt_path.read_text(encoding="utf-8", errors="ignore").splitlines()
    text_lines = []
    seen = set()
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if line.startswith("WEBVTT") or line.startswith("Kind:") or line.startswith("Language:"):
            continue
        if "-->" in line:  # timestamp line
            continue
        if re.match(r"^\d+$", line):  # cue index number
            continue
        # strip inline tags like <00:00:01.234><c> ... </c>
        clean = re.sub(r"<[^>]+>", "", line).strip()
        if not clean:
            continue
        if clean not in seen:  # auto-captions repeat lines across cues a lot
            text_lines.append(clean)
            seen.add(clean)
    return " ".join(text_lines)


def download_captions(video_id: str, out_stub: Path) -> tuple[Path | None, bool]:
    """
    Use yt-dlp to fetch captions only (no video/audio download).
    Tries manual captions first, falls back to auto-generated ones.
    Returns (path_to_vtt_or_None, was_rate_limited).
    """
    url = f"https://www.youtube.com/watch?v={video_id}"
    cmd = [
        sys.executable, "-m", "yt_dlp",
        "--skip-download",
        "--extractor-args", "youtube:player_client=tv,android",
        "--write-subs",
        "--write-auto-sub",
        "--sub-lang", LANG,
        "--sub-format", "vtt",
        "--output", str(out_stub) + ".%(ext)s",
    ]
    if COOKIES_FROM_BROWSER:
        cmd += ["--cookies-from-browser", COOKIES_FROM_BROWSER]
    cmd.append(url)

    for attempt in range(1, MAX_RETRIES + 1):
        try:
            result = subprocess.run(
                cmd, capture_output=True, text=True, timeout=YTDLP_TIMEOUT
            )
        except subprocess.TimeoutExpired:
            print(f"  [!] timeout fetching captions for {video_id}")
            return None, False

        if result.returncode == 0:
            break  # success, stop retrying

        stderr = result.stderr.strip()
        if "429" in stderr or "Too Many Requests" in stderr:
            wait = RATE_LIMIT_BACKOFF * attempt  # back off more each retry
            print(f"  [!] rate limited on {video_id}, waiting {wait}s "
                  f"(attempt {attempt}/{MAX_RETRIES}) ...")
            time.sleep(wait)
            continue  # retry
        else:
            print(f"  [!] yt-dlp error for {video_id}: {stderr[:500]}")
            return None, False  # non-rate-limit error, don't retry
    else:
        print(f"  [!] gave up on {video_id} after {MAX_RETRIES} rate-limit retries")
        return None, True  # signal: this failure was rate-limit related

    # yt-dlp names the file <stub>.<lang>.vtt
    candidate = out_stub.with_suffix(f".{LANG}.vtt")
    if candidate.exists():
        return candidate, False

    # sometimes it's just <stub>.vtt
    fallback = out_stub.with_suffix(".vtt")
    if fallback.exists():
        return fallback, False

    return None, False


def main():
    if len(sys.argv) != 2:
        print("Usage: python collect_transcripts.py <tracker.xlsx>")
        sys.exit(1)

    xlsx_path = Path(sys.argv[1])
    if not xlsx_path.exists():
        print(f"File not found: {xlsx_path}")
        sys.exit(1)

    TRANSCRIPTS_DIR.mkdir(exist_ok=True)

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb["videos"]

    headers = [cell.value for cell in ws[1]]
    col = {name: idx + 1 for idx, name in enumerate(headers)}

    required = ["motorcycle_id", "youtube_url", "youtube_video_id",
                "transcript_status", "transcript_path"]
    missing = [c for c in required if c not in col]
    if missing:
        print(f"Missing expected columns in 'videos' sheet: {missing}")
        sys.exit(1)

    total = 0
    done = 0
    skipped = 0
    failed = 0
    processed_since_batch = 0
    consecutive_rate_limit_fails = 0

    for row in range(2, ws.max_row + 1):
        mid = ws.cell(row=row, column=col["motorcycle_id"]).value
        url = ws.cell(row=row, column=col["youtube_url"]).value
        vid_id_cell = ws.cell(row=row, column=col["youtube_video_id"]).value
        status_cell = ws.cell(row=row, column=col["transcript_status"])
        path_cell = ws.cell(row=row, column=col["transcript_path"])

        if not url:
            continue  # no video collected for this motorcycle yet

        total += 1

        # skip if already done
        if status_cell.value == "done" and path_cell.value:
            skipped += 1
            continue

        video_id = extract_video_id(url)
        if not video_id and vid_id_cell and re.fullmatch(r"[A-Za-z0-9_-]{11}", str(vid_id_cell).strip()):
            video_id = str(vid_id_cell).strip()  # only trust pasted ID if it's a valid 11-char id
        if not video_id:
            print(f"[{mid}] could not parse video id from '{url}'")
            status_cell.value = "failed"
            failed += 1
            continue

        print(f"[{mid}] fetching captions for {video_id} ...")
        out_stub = TRANSCRIPTS_DIR / f"{mid}"
        vtt_file, was_rate_limited = download_captions(video_id, out_stub)

        if not vtt_file:
            status_cell.value = "failed"
            failed += 1
            if was_rate_limited:
                consecutive_rate_limit_fails += 1
                if consecutive_rate_limit_fails >= CONSECUTIVE_FAIL_LIMIT:
                    print(f"  [!] {consecutive_rate_limit_fails} rate-limit failures in a "
                          f"row — cooling down for {COOLDOWN_SECONDS}s before continuing ...")
                    wb.save(xlsx_path)  # save progress before the long pause
                    time.sleep(COOLDOWN_SECONDS)
                    consecutive_rate_limit_fails = 0
            time.sleep(SLEEP_BETWEEN_VIDEOS)
            continue

        consecutive_rate_limit_fails = 0  # reset streak on any success

        text = vtt_to_plain_text(vtt_file)
        txt_path = TRANSCRIPTS_DIR / f"{mid}.txt"
        txt_path.write_text(text, encoding="utf-8")
        vtt_file.unlink(missing_ok=True)  # clean up the raw vtt

        status_cell.value = "done"
        path_cell.value = str(txt_path)
        done += 1

        # save progress periodically so a crash doesn't lose everything
        if done % 10 == 0:
            wb.save(xlsx_path)
            print(f"  ... progress saved ({done} done so far)")

        processed_since_batch += 1
        if processed_since_batch >= BATCH_SIZE:
            print(f"  ... pausing {BATCH_PAUSE}s after a batch of {BATCH_SIZE} to ease off YouTube ...")
            wb.save(xlsx_path)
            time.sleep(BATCH_PAUSE)
            processed_since_batch = 0
        else:
            time.sleep(SLEEP_BETWEEN_VIDEOS)

    wb.save(xlsx_path)

    print("\n--- Summary ---")
    print(f"Total videos with a URL: {total}")
    print(f"Transcripts collected:   {done}")
    print(f"Already done (skipped):  {skipped}")
    print(f"Failed:                  {failed}")
    print(f"Transcripts saved to:    {TRANSCRIPTS_DIR.resolve()}")
    print(f"Workbook updated:        {xlsx_path.resolve()}")


if __name__ == "__main__":
    main()